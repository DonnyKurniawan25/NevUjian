import io
from decimal import Decimal
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import viewsets, generics, permissions, status
from rest_framework.decorators import action
from rest_framework.response import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import re
import json
import requests
from accounts.views import IsAdminOrTeacher
from .models import (
    Exam, Question, AnswerChoice, ExamSession, StudentAnswer, GuestParticipant, AISetting,
)
from .serializers import (
    ExamListSerializer, ExamDetailSerializer,
    QuestionSerializer, QuestionStudentSerializer,
    ExamSessionSerializer, StudentAnswerSerializer,
    ExamResultSerializer, GuestParticipantSerializer, AISettingSerializer,
)
from .utils import calculate_score, get_client_ip


class ExamViewSet(viewsets.ModelViewSet):
    """CRUD for exams - admin/teacher only."""
    permission_classes = [IsAdminOrTeacher]

    def get_serializer_class(self):
        if self.action == 'list':
            return ExamListSerializer
        return ExamDetailSerializer

    def get_queryset(self):
        queryset = Exam.objects.all()
        user = self.request.user

        # Teachers only see their own exams
        if user.role == 'teacher':
            queryset = queryset.filter(teacher=user)

        # Filters
        subject = self.request.query_params.get('subject')
        class_target = self.request.query_params.get('class_target')
        access_mode = self.request.query_params.get('access_mode')
        is_active = self.request.query_params.get('is_active')
        search = self.request.query_params.get('search')

        if subject:
            queryset = queryset.filter(subject__icontains=subject)
        if class_target:
            queryset = queryset.filter(class_target__icontains=class_target)
        if access_mode:
            queryset = queryset.filter(access_mode=access_mode)
        if is_active is not None:
            queryset = queryset.filter(is_active=is_active.lower() == 'true')
        if search:
            queryset = queryset.filter(title__icontains=search)

        return queryset

    def perform_create(self, serializer):
        serializer.save(teacher=self.request.user)

    @action(detail=True, methods=['post'])
    def toggle_active(self, request, pk=None):
        """Toggle exam active status."""
        exam = self.get_object()
        exam.is_active = not exam.is_active
        exam.save(update_fields=['is_active'])
        return Response({'is_active': exam.is_active})

    @action(detail=True, methods=['get'])
    def results(self, request, pk=None):
        """Get unified results for both login and guest students."""
        exam = self.get_object()
        sessions = ExamSession.objects.filter(exam=exam).select_related(
            'student', 'student__student_profile', 'guest_participant'
        ).order_by('-started_at')

        # Filter by participant type
        p_type = request.query_params.get('type')
        if p_type == 'login':
            sessions = sessions.filter(student__isnull=False)
        elif p_type == 'guest':
            sessions = sessions.filter(guest_participant__isnull=False)

        # Filter by status
        p_status = request.query_params.get('status')
        if p_status:
            sessions = sessions.filter(status=p_status)

        serializer = ExamResultSerializer(sessions, many=True)
        return Response({
            'exam': ExamListSerializer(exam).data,
            'results': serializer.data,
            'summary': {
                'total_participants': sessions.count(),
                'login_participants': sessions.filter(student__isnull=False).count(),
                'guest_participants': sessions.filter(guest_participant__isnull=False).count(),
                'submitted': sessions.filter(status='submitted').count(),
                'in_progress': sessions.filter(status='in_progress').count(),
                'average_score': self._calc_avg(sessions),
                'highest_score': self._calc_max(sessions),
                'lowest_score': self._calc_min(sessions),
                'pass_count': sessions.filter(
                    score__gte=exam.passing_score, status='submitted'
                ).count(),
            }
        })

    def _calc_avg(self, sessions):
        scored = sessions.filter(score__isnull=False)
        if not scored.exists():
            return 0
        from django.db.models import Avg
        return float(scored.aggregate(avg=Avg('score'))['avg'] or 0)

    def _calc_max(self, sessions):
        scored = sessions.filter(score__isnull=False)
        if not scored.exists():
            return 0
        from django.db.models import Max
        return float(scored.aggregate(mx=Max('score'))['mx'] or 0)

    def _calc_min(self, sessions):
        scored = sessions.filter(score__isnull=False)
        if not scored.exists():
            return 0
        from django.db.models import Min
        return float(scored.aggregate(mn=Min('score'))['mn'] or 0)


class QuestionViewSet(viewsets.ModelViewSet):
    """CRUD for questions within an exam."""
    serializer_class = QuestionSerializer
    permission_classes = [IsAdminOrTeacher]

    def get_queryset(self):
        exam_id = self.kwargs.get('exam_id')
        return Question.objects.filter(exam_id=exam_id).prefetch_related('choices')

    def perform_create(self, serializer):
        exam_id = self.kwargs.get('exam_id')
        exam = Exam.objects.get(pk=exam_id)
        order = exam.questions.count() + 1
        serializer.save(exam=exam, order=order)

    @action(detail=False, methods=['post'])
    def reorder(self, request, exam_id=None):
        """Reorder questions."""
        order_data = request.data.get('order', [])
        for item in order_data:
            Question.objects.filter(
                id=item['id'], exam_id=exam_id
            ).update(order=item['order'])
        return Response({'detail': 'Urutan soal berhasil diperbarui.'})

    @action(detail=False, methods=['post'])
    def bulk_create(self, request, exam_id=None):
        """Bulk create questions."""
        exam = Exam.objects.get(pk=exam_id)
        questions_data = request.data.get('questions', [])
        created = []
        start_order = exam.questions.count() + 1

        for i, q_data in enumerate(questions_data):
            choices_data = q_data.pop('choices', [])
            question = Question.objects.create(
                exam=exam, order=start_order + i, **q_data
            )
            for j, c_data in enumerate(choices_data):
                AnswerChoice.objects.create(question=question, order=j, **c_data)
            created.append(question)

        serializer = QuestionSerializer(created, many=True)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['post'], url_path='import-excel')
    def import_excel(self, request, exam_id=None):
        """Import questions from an Excel spreadsheet."""
        from openpyxl import load_workbook
        from django.db import transaction

        exam = Exam.objects.get(pk=exam_id)
        excel_file = request.FILES.get('file')

        if not excel_file:
            return Response({'error': 'File Excel tidak ditemukan.'}, status=status.HTTP_400_BAD_REQUEST)

        # Validate file extension
        if not excel_file.name.endswith(('.xlsx', '.xls')):
            return Response({'error': 'Format berkas harus berupa Excel (.xlsx atau .xls).'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = load_workbook(excel_file, data_only=True)
            ws = wb.active
        except Exception as e:
            return Response({'error': f'Gagal membaca file Excel: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        # Parse rows
        questions_to_create = []
        errors = []
        
        # We start reading from row 2 (assuming row 1 is headers)
        row_idx = 2
        
        while True:
            # Check if row is empty
            # We look at Pertanyaan (col 2) and Tipe Soal (col 3)
            pertanyaan_cell = ws.cell(row=row_idx, column=2).value
            tipe_cell = ws.cell(row=row_idx, column=3).value
            
            if pertanyaan_cell is None and tipe_cell is None:
                # Check if we should break.
                next_empty = True
                for offset in range(1, 4):
                    if ws.cell(row=row_idx + offset, column=2).value is not None or ws.cell(row=row_idx + offset, column=3).value is not None:
                        next_empty = False
                        break
                if next_empty:
                    break
                else:
                    row_idx += 1
                    continue

            pertanyaan = str(pertanyaan_cell).strip() if pertanyaan_cell is not None else ""
            tipe_soal_raw = str(tipe_cell).strip().lower() if tipe_cell is not None else ""
            poin_val = ws.cell(row=row_idx, column=4).value
            
            # Validasi dasar
            if not pertanyaan:
                errors.append(f"Baris {row_idx}: Kolom 'Pertanyaan' tidak boleh kosong.")
                row_idx += 1
                continue

            # Determine type
            if tipe_soal_raw in ('pilihan ganda', 'pg', 'mc', 'multiple_choice', 'pilihan_ganda'):
                q_type = 'multiple_choice'
            elif tipe_soal_raw in ('essay', 'uraian', 'esay'):
                q_type = 'essay'
            else:
                q_type = 'multiple_choice'

            # Parse points
            try:
                points = int(poin_val) if poin_val is not None else 1
                if points < 0:
                    points = 1
            except ValueError:
                points = 1

            choices = []
            explanation = ""

            if q_type == 'multiple_choice':
                choice_a = ws.cell(row=row_idx, column=5).value
                choice_b = ws.cell(row=row_idx, column=6).value
                choice_c = ws.cell(row=row_idx, column=7).value
                choice_d = ws.cell(row=row_idx, column=8).value
                choice_e = ws.cell(row=row_idx, column=9).value
                kunci = ws.cell(row=row_idx, column=10).value
                
                # Check required choices
                if choice_a is None or choice_b is None or choice_c is None or choice_d is None:
                    errors.append(f"Baris {row_idx}: Pilihan jawaban A, B, C, dan D wajib diisi untuk Pilihan Ganda.")
                    row_idx += 1
                    continue

                if kunci is None:
                    errors.append(f"Baris {row_idx}: Kunci jawaban wajib diisi untuk Pilihan Ganda.")
                    row_idx += 1
                    continue

                kunci = str(kunci).strip().upper()
                if kunci not in ('A', 'B', 'C', 'D', 'E'):
                    errors.append(f"Baris {row_idx}: Kunci jawaban harus berupa salah satu huruf: A, B, C, D, atau E.")
                    row_idx += 1
                    continue

                if kunci == 'E' and choice_e is None:
                    errors.append(f"Baris {row_idx}: Kunci jawaban disetel ke 'E', tetapi Pilihan E kosong.")
                    row_idx += 1
                    continue

                # Add choices data
                choices.append({'text': str(choice_a).strip(), 'is_correct': kunci == 'A', 'order': 1})
                choices.append({'text': str(choice_b).strip(), 'is_correct': kunci == 'B', 'order': 2})
                choices.append({'text': str(choice_c).strip(), 'is_correct': kunci == 'C', 'order': 3})
                choices.append({'text': str(choice_d).strip(), 'is_correct': kunci == 'D', 'order': 4})
                if choice_e is not None and str(choice_e).strip() != "":
                    choices.append({'text': str(choice_e).strip(), 'is_correct': kunci == 'E', 'order': 5})

            explanation_cell = ws.cell(row=row_idx, column=11).value
            explanation = str(explanation_cell).strip() if explanation_cell is not None else ""

            questions_to_create.append({
                'question_text': pertanyaan,
                'question_type': q_type,
                'points': points,
                'choices': choices,
                'explanation': explanation,
            })
            
            row_idx += 1

        if errors:
            return Response({'errors': errors}, status=status.HTTP_400_BAD_REQUEST)

        if not questions_to_create:
            return Response({'error': 'Tidak ada data soal yang ditemukan dalam file Excel.'}, status=status.HTTP_400_BAD_REQUEST)

        # Database Transaction
        try:
            with transaction.atomic():
                start_order = exam.questions.count() + 1
                created_questions = []
                for idx, q_data in enumerate(questions_to_create):
                    choices_data = q_data.pop('choices', [])
                    question = Question.objects.create(
                        exam=exam,
                        order=start_order + idx,
                        **q_data
                    )
                    for c_data in choices_data:
                        AnswerChoice.objects.create(
                            question=question,
                            choice_text=c_data['text'],
                            is_correct=c_data['is_correct'],
                            order=c_data['order']
                        )
                    created_questions.append(question)
        except Exception as e:
            return Response({'error': f'Gagal menyimpan data ke database: {str(e)}'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        serializer = QuestionSerializer(created_questions, many=True)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'], url_path='download-template')
    def download_template(self, request, exam_id=None):
        """Download Excel template for importing questions."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        import io
        from django.http import HttpResponse

        wb = Workbook()
        ws = wb.active
        ws.title = 'Template Soal'

        # Headers
        headers = [
            'No', 'Pertanyaan', 'Tipe Soal', 'Poin',
            'Pilihan A', 'Pilihan B', 'Pilihan C', 'Pilihan D', 'Pilihan E',
            'Kunci Jawaban', 'Pembahasan'
        ]

        # Style headers
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        # Sample data
        samples = [
            [
                1, 'Berapakah hasil dari 5 + 3?', 'Pilihan Ganda', 10,
                '6', '8', '10', '12', '', 'B', 'Karena 5 + 3 = 8.'
            ],
            [
                2, 'Jelaskan perbedaan antara RAM dan Harddisk!', 'Essay', 15,
                '', '', '', '', '', '', 'RAM bersifat volatile sedangkan Harddisk non-volatile.'
            ]
        ]

        for row_idx, sample in enumerate(samples, 2):
            for col_idx, val in enumerate(sample, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx in (1, 3, 4, 10):
                    cell.alignment = Alignment(horizontal='center')
                else:
                    cell.alignment = Alignment(horizontal='left')

        # Auto width
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="template_soal_nevujian.xlsx"'
        return response

    @action(detail=False, methods=['delete'], url_path='delete-all')
    def delete_all(self, request, exam_id=None):
        """Delete all questions for this exam."""
        exam = generics.get_object_or_404(Exam, pk=exam_id)
        
        from django.db import transaction
        try:
            with transaction.atomic():
                exam.questions.all().delete()
            return Response({'detail': 'Semua soal berhasil dihapus.'}, status=status.HTTP_200_OK)
        except Exception as e:
            return Response({'error': f"Gagal menghapus semua soal: {str(e)}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class StudentExamListView(generics.ListAPIView):
    """List available exams for logged-in students."""
    serializer_class = ExamListSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        queryset = Exam.objects.filter(
            is_active=True,
            access_mode__in=['login_required', 'both'],
        )

        # Filter by student's class
        if hasattr(user, 'student_profile'):
            student_class = user.student_profile.class_name
            queryset = queryset.filter(class_target__icontains=student_class)

        return queryset


class StudentExamSessionView(generics.GenericAPIView):
    """Start or resume an exam session for logged-in student."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, exam_id):
        """Start a new exam session."""
        exam = Exam.objects.get(pk=exam_id)

        # Validate
        if not exam.is_active:
            return Response({'error': 'Ujian tidak aktif.'}, status=400)
        if not exam.is_ongoing:
            if exam.is_upcoming:
                return Response({'error': 'Ujian belum dimulai.'}, status=400)
            return Response({'error': 'Ujian telah berakhir.'}, status=400)
        if exam.access_mode == 'guest_allowed':
            return Response({'error': 'Ujian ini hanya untuk peserta tanpa login.'}, status=400)

        # Check if student already has an active session
        existing = ExamSession.objects.filter(
            exam=exam, student=request.user, status='in_progress'
        ).first()
        if existing:
            serializer = ExamSessionSerializer(existing)
            return Response(serializer.data)

        # Check if already submitted
        submitted = ExamSession.objects.filter(
            exam=exam, student=request.user, status__in=['submitted', 'timed_out']
        ).exists()
        if submitted:
            return Response({'error': 'Anda sudah mengerjakan ujian ini.'}, status=400)

        # Validate token if required
        if exam.require_exam_token:
            token = request.data.get('exam_token')
            if token != exam.exam_token:
                return Response({'error': 'Token ujian salah.'}, status=400)

        # Create session
        end_time = min(
            timezone.now() + timezone.timedelta(minutes=exam.duration_minutes),
            exam.end_time
        )
        session = ExamSession.objects.create(
            exam=exam,
            student=request.user,
            end_time=end_time,
            ip_address=get_client_ip(request),
            user_agent=request.META.get('HTTP_USER_AGENT', ''),
        )

        serializer = ExamSessionSerializer(session)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class StudentExamSessionDetailView(generics.GenericAPIView):
    """Get session details, submit answers, submit exam, record violations."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, session_id):
        """Get session with questions."""
        session = ExamSession.objects.select_related(
            'exam'
        ).get(pk=session_id, student=request.user)

        # Auto-submit if timed out
        if session.is_timed_out:
            session.status = 'timed_out'
            session.submitted_at = session.end_time
            session.score = calculate_score(session)
            session.save()

        pg_questions = session.exam.questions.prefetch_related('choices').filter(question_type='multiple_choice')
        essay_questions = session.exam.questions.prefetch_related('choices').filter(question_type='essay')
        
        if session.exam.shuffle_questions:
            pg_questions = pg_questions.order_by('?')
            essay_questions = essay_questions.order_by('?')
        else:
            pg_questions = pg_questions.order_by('order')
            essay_questions = essay_questions.order_by('order')
            
        questions = list(pg_questions) + list(essay_questions)

        session_data = ExamSessionSerializer(session).data
        session_data['questions'] = QuestionStudentSerializer(questions, many=True).data

        # Include existing answers
        existing_answers = StudentAnswer.objects.filter(
            session=session
        ).values('question_id', 'selected_choice_id', 'essay_answer')
        session_data['saved_answers'] = {
            str(a['question_id']): {
                'selected_choice': a['selected_choice_id'],
                'essay_answer': a['essay_answer'],
            } for a in existing_answers
        }

        return Response(session_data)


class AnswerView(generics.GenericAPIView):
    """Save an answer for a question."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, session_id):
        session = ExamSession.objects.get(pk=session_id, student=request.user)

        if session.status != 'in_progress':
            return Response({'error': 'Sesi ujian sudah berakhir.'}, status=400)
        if session.is_timed_out:
            return Response({'error': 'Waktu ujian telah habis.'}, status=400)

        question_id = request.data.get('question_id')
        selected_choice_id = request.data.get('selected_choice')
        essay_answer = request.data.get('essay_answer', '')

        answer, created = StudentAnswer.objects.update_or_create(
            session=session,
            question_id=question_id,
            defaults={
                'selected_choice_id': selected_choice_id,
                'essay_answer': essay_answer,
            }
        )

        return Response(StudentAnswerSerializer(answer).data)


class SubmitExamView(generics.GenericAPIView):
    """Submit the entire exam."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, session_id):
        session = ExamSession.objects.get(pk=session_id, student=request.user)

        if session.status != 'in_progress':
            return Response({'error': 'Sesi ujian sudah berakhir.'}, status=400)

        session.status = 'submitted'
        session.submitted_at = timezone.now()
        session.score = calculate_score(session)
        session.save()

        return Response(ExamSessionSerializer(session).data)


class ViolationView(generics.GenericAPIView):
    """Record exam violation."""
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, session_id):
        session = ExamSession.objects.get(pk=session_id, student=request.user)

        if session.status != 'in_progress':
            return Response({'error': 'Sesi ujian sudah berakhir.'}, status=400)

        session.violation_count += 1
        if session.violation_count >= session.exam.max_violations:
            session.status = 'terminated'
            session.submitted_at = timezone.now()
            session.score = calculate_score(session)

        session.save()

        return Response({
            'violation_count': session.violation_count,
            'max_violations': session.exam.max_violations,
            'terminated': session.status == 'terminated',
        })


class ExportExcelView(generics.GenericAPIView):
    """Export exam results to Excel."""
    permission_classes = [IsAdminOrTeacher]

    def get(self, request, exam_id):
        exam = Exam.objects.get(pk=exam_id)
        sessions = ExamSession.objects.filter(exam=exam).select_related(
            'student', 'student__student_profile', 'guest_participant'
        ).order_by('started_at')

        wb = Workbook()
        ws = wb.active
        ws.title = 'Hasil Ujian'

        # Styles
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )

        # Title
        ws.merge_cells('A1:J1')
        ws['A1'] = f'Hasil Ujian: {exam.title}'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f'Mata Pelajaran: {exam.subject} | Kelas: {exam.class_target}'
        ws['A3'] = f'Tanggal: {exam.start_time.strftime("%d/%m/%Y %H:%M")} - {exam.end_time.strftime("%d/%m/%Y %H:%M")}'
        ws['A4'] = f'Mode Akses: {exam.get_access_mode_display()}'

        # Headers
        headers = [
            'No', 'Nama Peserta', 'NIS', 'NISN', 'Kelas',
            'Tipe Peserta', 'Nilai', 'Status', 'Waktu Mulai',
            'Waktu Selesai', 'Pelanggaran'
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        # Data
        for row_num, session in enumerate(sessions, 7):
            if session.student:
                name = session.student.get_full_name() or session.student.username
                nis = getattr(session.student, 'student_profile', None)
                nis_val = nis.nis if nis else '-'
                nisn_val = nis.nisn if nis else '-'
                class_val = nis.class_name if nis else '-'
                p_type = 'Siswa Login'
            elif session.guest_participant:
                name = session.guest_participant.full_name
                nis_val = session.guest_participant.nis
                nisn_val = session.guest_participant.nisn or '-'
                class_val = session.guest_participant.class_name
                p_type = 'Peserta Tanpa Login'
            else:
                name = '-'
                nis_val = nisn_val = class_val = '-'
                p_type = '-'

            data = [
                row_num - 6,
                name,
                nis_val,
                nisn_val,
                class_val,
                p_type,
                float(session.score) if session.score else 0,
                session.get_status_display(),
                session.started_at.strftime('%d/%m/%Y %H:%M') if session.started_at else '-',
                session.submitted_at.strftime('%d/%m/%Y %H:%M') if session.submitted_at else '-',
                session.violation_count,
            ]

            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row_num, column=col, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')

        # Auto-width
        from openpyxl.utils import get_column_letter
        for col in ws.columns:
            max_length = 0
            column_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = min(max_length + 4, 40)

        # Summary sheet
        ws2 = wb.create_sheet('Ringkasan')
        ws2['A1'] = 'Ringkasan Hasil'
        ws2['A1'].font = Font(bold=True, size=14)

        summary_data = [
            ('Total Peserta', sessions.count()),
            ('Siswa Login', sessions.filter(student__isnull=False).count()),
            ('Peserta Guest', sessions.filter(guest_participant__isnull=False).count()),
            ('Sudah Submit', sessions.filter(status='submitted').count()),
            ('Sedang Berlangsung', sessions.filter(status='in_progress').count()),
            ('Waktu Habis', sessions.filter(status='timed_out').count()),
            ('Dihentikan', sessions.filter(status='terminated').count()),
        ]

        for i, (label, value) in enumerate(summary_data, 3):
            ws2.cell(row=i, column=1, value=label).font = Font(bold=True)
            ws2.cell(row=i, column=2, value=value)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'hasil_ujian_{exam.exam_code}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class ExportPDFView(generics.GenericAPIView):
    """Export exam results to PDF."""
    permission_classes = [IsAdminOrTeacher]

    def get(self, request, exam_id):
        from xhtml2pdf import pisa

        exam = Exam.objects.get(pk=exam_id)
        sessions = ExamSession.objects.filter(exam=exam).select_related(
            'student', 'student__student_profile', 'guest_participant'
        ).order_by('started_at')

        rows = ''
        for i, session in enumerate(sessions, 1):
            if session.student:
                name = session.student.get_full_name() or session.student.username
                profile = getattr(session.student, 'student_profile', None)
                nis = profile.nis if profile else '-'
                nisn = profile.nisn if profile else '-'
                kelas = profile.class_name if profile else '-'
                tipe = 'Login'
            elif session.guest_participant:
                name = session.guest_participant.full_name
                nis = session.guest_participant.nis
                nisn = session.guest_participant.nisn or '-'
                kelas = session.guest_participant.class_name
                tipe = 'Guest'
            else:
                name = nis = nisn = kelas = tipe = '-'

            score = float(session.score) if session.score else 0
            status_text = session.get_status_display()
            started = session.started_at.strftime('%d/%m/%Y %H:%M') if session.started_at else '-'
            submitted = session.submitted_at.strftime('%d/%m/%Y %H:%M') if session.submitted_at else '-'

            rows += f'''<tr>
                <td>{i}</td>
                <td>{name}</td>
                <td>{nis}</td>
                <td>{nisn}</td>
                <td>{kelas}</td>
                <td>{tipe}</td>
                <td>{score:.1f}</td>
                <td>{status_text}</td>
                <td>{started}</td>
                <td>{submitted}</td>
                <td>{session.violation_count}</td>
            </tr>'''

        html = f'''
        <html>
        <head>
            <style>
                body {{ font-family: Arial, sans-serif; font-size: 10px; }}
                h1 {{ font-size: 16px; color: #1F4E79; }}
                h3 {{ font-size: 12px; color: #333; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 10px; }}
                th {{ background-color: #1F4E79; color: white; padding: 6px; text-align: center; font-size: 9px; }}
                td {{ padding: 4px; text-align: center; border: 1px solid #ddd; font-size: 9px; }}
                tr:nth-child(even) {{ background-color: #f2f2f2; }}
                .info {{ margin-bottom: 5px; }}
            </style>
        </head>
        <body>
            <h1>Hasil Ujian: {exam.title}</h1>
            <p class="info">Mata Pelajaran: {exam.subject} | Kelas: {exam.class_target}</p>
            <p class="info">Waktu: {exam.start_time.strftime("%d/%m/%Y %H:%M")} - {exam.end_time.strftime("%d/%m/%Y %H:%M")}</p>
            <p class="info">Mode Akses: {exam.get_access_mode_display()} | Total Peserta: {sessions.count()}</p>
            <table>
                <tr>
                    <th>No</th><th>Nama</th><th>NIS</th><th>NISN</th>
                    <th>Kelas</th><th>Tipe</th><th>Nilai</th><th>Status</th>
                    <th>Mulai</th><th>Selesai</th><th>Pelanggaran</th>
                </tr>
                {rows}
            </table>
        </body>
        </html>
        '''

        response = HttpResponse(content_type='application/pdf')
        filename = f'hasil_ujian_{exam.exam_code}.pdf'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        pisa.CreatePDF(html, dest=response)
        return response


class MyExamSessionsView(generics.ListAPIView):
    """List all exam sessions for the current student."""
    serializer_class = ExamSessionSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        return ExamSession.objects.filter(
            student=self.request.user
        ).select_related('exam', 'student', 'guest_participant').order_by('-started_at')


def clean_and_parse_json(response_text):
    text = response_text.strip()
    if text.startswith("```"):
        match = re.match(r"^```(?:json)?\s*(.*?)\s*```$", text, re.DOTALL | re.IGNORECASE)
        if match:
            text = match.group(1).strip()
    return json.loads(text)


class AISettingView(generics.RetrieveUpdateAPIView):
    """Retrieve and update AI configuration settings."""
    serializer_class = AISettingSerializer
    permission_classes = [IsAdminOrTeacher]

    def get_object(self):
        obj, created = AISetting.objects.get_or_create(pk=1)
        return obj

    def perform_update(self, serializer):
        serializer.save(updated_by=self.request.user)


class AITestConnectionView(generics.GenericAPIView):
    """Test connection to the AI provider without saving it first (or test saved settings)."""
    permission_classes = [IsAdminOrTeacher]

    def post(self, request):
        provider = request.data.get('provider')
        base_url = request.data.get('base_url')
        model_name = request.data.get('model_name')
        api_key = request.data.get('api_key')
        timeout = request.data.get('timeout', 120)
        temperature = request.data.get('temperature', 0.4)
        max_tokens = request.data.get('max_tokens', 8000)

        saved_setting, _ = AISetting.objects.get_or_create(pk=1)

        # Merge with saved settings if fields are empty/omitted
        if api_key and '*' in api_key:
            api_key = saved_setting.api_key
        
        if not provider:
            provider = saved_setting.provider
        if not base_url:
            base_url = saved_setting.base_url
        if not model_name:
            model_name = saved_setting.model_name
        if api_key is None or api_key == '':
            api_key = saved_setting.api_key

        if not base_url:
            return Response({'error': 'Base URL wajib diisi.'}, status=status.HTTP_400_BAD_REQUEST)

        # Normalize base_url
        normalized_url = base_url.rstrip('/') + '/chat/completions'

        headers = {
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json"
        }

        payload = {
            "model": model_name,
            "messages": [
                {"role": "user", "content": "ping"}
            ],
            "temperature": 0.1,
            "max_tokens": 10
        }

        try:
            timeout_val = int(timeout)
        except ValueError:
            timeout_val = 120

        is_success = False
        response_text = None
        error_msg = None

        try:
            response = requests.post(normalized_url, json=payload, headers=headers, timeout=timeout_val)
            if response.status_code == 200:
                is_success = True
                try:
                    resp_json = response.json()
                    response_text = resp_json['choices'][0]['message']['content']
                except Exception:
                    response_text = response.text[:200]
            else:
                error_msg = f"HTTP {response.status_code}: {response.text[:500]}"
        except requests.exceptions.Timeout:
            error_msg = "Request Timeout (Waktu habis)"
        except Exception as e:
            error_msg = str(e)

        # Save connection status in database
        saved_setting.last_checked = timezone.now()
        if is_success:
            saved_setting.last_status = f"Terhubung. Response: {response_text}"
        else:
            saved_setting.last_status = f"Terputus. Error: {error_msg}"
        saved_setting.save()

        return Response({
            'success': is_success,
            'message': "Koneksi Berhasil" if is_success else "Koneksi Gagal",
            'response': response_text or error_msg,
            'last_checked': saved_setting.last_checked,
            'last_status': saved_setting.last_status
        })


class AIGenerateQuestionsView(generics.GenericAPIView):
    """Automatically generate exam questions using AI."""
    permission_classes = [IsAdminOrTeacher]

    def post(self, request, exam_id):
        exam = generics.get_object_or_404(Exam, pk=exam_id)
        settings_obj = AISetting.objects.filter(pk=1).first()
        
        if not settings_obj or not settings_obj.is_active:
            return Response({'error': 'Integrasi AI dinonaktifkan atau belum dikonfigurasi.'}, status=status.HTTP_400_BAD_REQUEST)
        if not settings_obj.api_key:
            return Response({'error': 'API Key AI belum diatur.'}, status=status.HTTP_400_BAD_REQUEST)

        topic = request.data.get('topic', 'Umum')
        question_type = request.data.get('question_type', 'multiple_choice')
        try:
            count = int(request.data.get('count', 5))
            if count <= 0:
                count = 5
            elif count > 50:
                count = 50
        except ValueError:
            count = 5
            
        try:
            points = int(request.data.get('points', 10))
            if points < 0:
                points = 10
        except ValueError:
            points = 10

        import math
        batch_size = 10
        total_requested = count
        num_batches = math.ceil(total_requested / batch_size)
        
        questions_data = []
        errors = []

        base_url = settings_obj.base_url.rstrip('/')
        url = f"{base_url}/chat/completions"

        headers = {
            "Authorization": f"Bearer {settings_obj.api_key}",
            "Content-Type": "application/json"
        }

        # Keep track of previously generated questions to avoid duplicates in subsequent prompts
        previously_generated_titles = []

        for i in range(num_batches):
            current_batch_count = min(batch_size, total_requested - i * batch_size)
            
            avoid_dupes_instruction = ""
            if previously_generated_titles:
                avoid_dupes_instruction = (
                    f"PENTING: Jangan buat soal yang mirip dengan topik-topik soal berikut: "
                    f"{', '.join(previously_generated_titles[:15])}. Hasilkan soal yang unik."
                )

            if question_type == 'multiple_choice':
                user_prompt = (
                    f"Buatkan {current_batch_count} soal pilihan ganda tentang topik '{topic}'.\n"
                    f"Setiap soal harus memiliki minimal 4 pilihan jawaban (A, B, C, D) atau maksimal 5 (A, B, C, D, E), "
                    f"dengan tepat satu pilihan jawaban yang benar (is_correct=true, sisanya false).\n"
                    f"Sediakan juga pembahasan singkat di field 'explanation'.\n"
                    f"{avoid_dupes_instruction}\n\n"
                    f"Format output JSON yang wajib diikuti:\n"
                    f"[\n"
                    f"  {{\n"
                    f"    \"question_text\": \"Teks pertanyaan soal PG\",\n"
                    f"    \"choices\": [\n"
                    f"      {{\"choice_text\": \"Pilihan A\", \"is_correct\": true}},\n"
                    f"      {{\"choice_text\": \"Pilihan B\", \"is_correct\": false}},\n"
                    f"      {{\"choice_text\": \"Pilihan C\", \"is_correct\": false}},\n"
                    f"      {{\"choice_text\": \"Pilihan D\", \"is_correct\": false}}\n"
                    f"    ],\n"
                    f"    \"explanation\": \"Pembahasan soal...\"\n"
                    f"  }}\n"
                    f"]"
                )
            else:
                user_prompt = (
                    f"Buatkan {current_batch_count} soal essay tentang topik '{topic}'.\n"
                    f"Untuk soal essay, sediakan kunci jawaban/rubrik penilaian singkat di field 'explanation'.\n"
                    f"{avoid_dupes_instruction}\n\n"
                    f"Format output JSON yang wajib diikuti:\n"
                    f"[\n"
                    f"  {{\n"
                    f"    \"question_text\": \"Teks pertanyaan soal essay\",\n"
                    f"    \"explanation\": \"Kunci jawaban / pembahasan soal essay...\"\n"
                    f"  }}\n"
                    f"]"
                )

            # Cap max_tokens in the payload to a safe maximum for generation (e.g. 4096)
            # as most providers do not support output tokens above 4096/8192 and will return HTTP 400.
            max_tokens_payload = min(settings_obj.max_tokens or 2048, 4096)

            payload = {
                "model": settings_obj.model_name,
                "messages": [
                    {"role": "system", "content": "Anda adalah asisten pembuat soal ujian profesional yang hanya menghasilkan JSON murni sesuai spesifikasi."},
                    {"role": "user", "content": user_prompt}
                ],
                "temperature": settings_obj.temperature,
                "max_tokens": max_tokens_payload
            }

            try:
                response = requests.post(url, json=payload, headers=headers, timeout=settings_obj.timeout)
                if response.status_code != 200:
                    errors.append(f"Batch {i+1} gagal (HTTP {response.status_code}): {response.text[:200]}")
                    continue
                
                resp_json = response.json()
                content = resp_json['choices'][0]['message']['content']
                
                batch_questions = clean_and_parse_json(content)
                if isinstance(batch_questions, list):
                    for q in batch_questions:
                        q_text = q.get('question_text', '').strip()
                        if q_text:
                            # Avoid duplicates from the model's own response within this generation run
                            if q_text not in [x.get('question_text') for x in questions_data]:
                                questions_data.append(q)
                                previously_generated_titles.append(q_text[:40])
                else:
                    errors.append(f"Batch {i+1} tidak mengembalikan array JSON.")
            except Exception as e:
                errors.append(f"Batch {i+1} error: {str(e)}")

        if not questions_data:
            return Response({
                'error': 'Gagal menghasilkan soal dari AI. AI tidak mengembalikan format JSON yang valid atau semua batch gagal.',
                'errors': errors
            }, status=status.HTTP_400_BAD_REQUEST)

        # Database Transaction
        from django.db import transaction
        created_questions = []
        try:
            with transaction.atomic():
                start_order = exam.questions.count() + 1
                for idx, q_data in enumerate(questions_data):
                    q_text = q_data.get('question_text', '').strip()
                    if not q_text:
                        continue
                    explanation = q_data.get('explanation', '').strip()
                    
                    question = Question.objects.create(
                        exam=exam,
                        question_text=q_text,
                        question_type=question_type,
                        points=points,
                        explanation=explanation,
                        order=start_order + idx
                    )
                    
                    if question_type == 'multiple_choice':
                        choices_list = q_data.get('choices', [])
                        for choice_idx, choice_data in enumerate(choices_list):
                            AnswerChoice.objects.create(
                                question=question,
                                choice_text=choice_data.get('choice_text', '').strip(),
                                is_correct=choice_data.get('is_correct', False),
                                order=choice_idx
                            )
                    
                    created_questions.append(question)
        except Exception as e:
            return Response({'error': f"Gagal menyimpan soal ke database: {str(e)}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        serializer = QuestionSerializer(created_questions, many=True)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class AIGradeEssaysView(generics.GenericAPIView):
    """Grade essay answers automatically using AI for a given exam session in a single batch request."""
    permission_classes = [IsAdminOrTeacher]

    def post(self, request, session_id):
        session = generics.get_object_or_404(ExamSession, pk=session_id)
        settings_obj = AISetting.objects.filter(pk=1).first()
        
        if not settings_obj or not settings_obj.is_active:
            return Response({'error': 'Integrasi AI dinonaktifkan atau belum dikonfigurasi.'}, status=status.HTTP_400_BAD_REQUEST)
        if not settings_obj.api_key:
            return Response({'error': 'API Key AI belum diatur.'}, status=status.HTTP_400_BAD_REQUEST)

        essay_answers = StudentAnswer.objects.filter(
            session=session,
            question__question_type='essay'
        ).select_related('question')

        if not essay_answers.exists():
            return Response({'detail': 'Tidak ada jawaban essay yang perlu dinilai.'})

        base_url = settings_obj.base_url.rstrip('/')
        url = f"{base_url}/chat/completions"
        headers = {
            "Authorization": f"Bearer {settings_obj.api_key}",
            "Content-Type": "application/json"
        }

        graded_count = 0
        errors = []
        prompt_data = []

        # Separate blank answers and construct batch payload for non-blank ones
        for answer in essay_answers:
            q = answer.question
            student_text = answer.essay_answer.strip()

            if not student_text:
                answer.points_earned = Decimal('0.00')
                answer.is_correct = False
                answer.ai_feedback = "Siswa tidak memberikan jawaban."
                answer.save()
                graded_count += 1
                continue

            prompt_data.append({
                "id": answer.id,
                "question": q.question_text,
                "key": q.explanation,
                "student_answer": student_text,
                "max_points": float(q.points)
            })

        if prompt_data:
            prompt = (
                f"Evaluasi semua jawaban essay siswa berikut berdasarkan kunci jawaban dan kriteria penilaian. "
                f"Untuk setiap jawaban, berikan nilai ('score') antara 0 hingga 'max_points' dan berikan masukan singkat ('feedback') dalam Bahasa Indonesia.\n"
                f"Format output harus berupa JSON array of objects dengan key: 'id' (integer), 'score' (number), dan 'feedback' (string).\n"
                f"Kembalikan HANYA format JSON murni tanpa markdown, penjelasan atau pembungkus lain.\n\n"
                f"Daftar Jawaban:\n"
                f"{json.dumps(prompt_data, ensure_ascii=False, indent=2)}\n\n"
                f"Contoh Response:\n"
                f"[\n"
                f"  {{\n"
                f"    \"id\": {prompt_data[0]['id']},\n"
                f"    \"score\": 7.5,\n"
                f"    \"feedback\": \"Penjelasan sudah baik, tetapi kurang detail pada bagian X.\"\n"
                f"  }}\n"
                f"]"
            )

            payload = {
                "model": settings_obj.model_name,
                "messages": [
                    {"role": "system", "content": "Anda adalah asisten koreksi ujian otomatis yang objektif dan mengembalikan format JSON array murni."},
                    {"role": "user", "content": prompt}
                ],
                "temperature": 0.2,
                "max_tokens": 2000
            }

            try:
                response = requests.post(url, json=payload, headers=headers, timeout=settings_obj.timeout)
                if response.status_code == 200:
                    try:
                        resp_json = response.json()
                    except Exception as json_err:
                        errors.append(f"Gagal membaca JSON dari provider AI (Status 200 OK): {str(json_err)}. Response mentah: {response.text[:200]}")
                        resp_json = None
                    
                    if resp_json:
                        try:
                            content = resp_json['choices'][0]['message']['content']
                        except (KeyError, IndexError) as key_err:
                            errors.append(f"Format respon AI tidak sesuai standard OpenAI: {str(key_err)}. Response: {str(resp_json)[:200]}")
                            content = None
                        
                        if content:
                            try:
                                ai_results = clean_and_parse_json(content)
                            except json.JSONDecodeError as parse_err:
                                errors.append(f"Gagal mengurai hasil AI sebagai JSON: {str(parse_err)}. Teks dari AI: {content[:200]}")
                                ai_results = None
                            
                            if ai_results:
                                # Parse results mapping
                                results_map = {}
                                if isinstance(ai_results, list):
                                    for item in ai_results:
                                        if isinstance(item, dict) and 'id' in item:
                                            results_map[item.get('id')] = item
                                elif isinstance(ai_results, dict):
                                    # Handle potential wrapping or object mapping
                                    for key, val in ai_results.items():
                                        if isinstance(val, dict):
                                            try:
                                                results_map[int(key)] = val
                                            except ValueError:
                                                pass
                                        elif key in ['grades', 'results', 'answers', 'data'] and isinstance(val, list):
                                            for item in val:
                                                if isinstance(item, dict) and 'id' in item:
                                                    results_map[item.get('id')] = item
                                            break

                                # Apply grades to database
                                for answer in essay_answers:
                                    if answer.id in results_map:
                                        item = results_map[answer.id]
                                        score_val = item.get('score', 0)
                                        feedback = item.get('feedback', '')
                                        
                                        try:
                                            score = Decimal(str(score_val))
                                        except Exception:
                                            score = Decimal('0.00')
                                            
                                        max_points = answer.question.points
                                        if score < 0:
                                            score = Decimal('0.00')
                                        elif score > max_points:
                                            score = Decimal(str(max_points))
                                            
                                        answer.points_earned = score
                                        answer.is_correct = score >= (Decimal(str(max_points)) / 2)
                                        answer.ai_feedback = feedback
                                        answer.save()
                                        graded_count += 1
                                    elif answer.essay_answer.strip():
                                        errors.append(f"Jawaban ID {answer.id} (soal {answer.question.order}) tidak ada dalam respon grading AI.")
                else:
                    errors.append(f"AI Provider returned HTTP {response.status_code}: {response.text[:200]}")
            except Exception as e:
                errors.append(f"Gagal menghubungi AI atau parsing data: {str(e)}")

        # Recalculate session score
        session.score = calculate_score(session)
        session.save()

        return Response({
            'success': len(errors) == 0,
            'graded_count': graded_count,
            'errors': errors,
            'new_score': session.score
        })


class AIGradeSingleAnswerView(generics.GenericAPIView):
    """Grade a single essay answer using AI."""
    permission_classes = [IsAdminOrTeacher]

    def post(self, request, answer_id):
        answer = generics.get_object_or_404(StudentAnswer, pk=answer_id)
        if answer.question.question_type != 'essay':
            return Response({'error': 'Jawaban ini bukan merupakan tipe soal essay.'}, status=status.HTTP_400_BAD_REQUEST)

        settings_obj = AISetting.objects.filter(pk=1).first()
        if not settings_obj or not settings_obj.is_active:
            return Response({'error': 'Integrasi AI dinonaktifkan atau belum dikonfigurasi.'}, status=status.HTTP_400_BAD_REQUEST)
        if not settings_obj.api_key:
            return Response({'error': 'API Key AI belum diatur.'}, status=status.HTTP_400_BAD_REQUEST)

        q = answer.question
        max_points = q.points
        student_text = answer.essay_answer.strip()

        if not student_text:
            answer.points_earned = Decimal('0.00')
            answer.is_correct = False
            answer.ai_feedback = "Siswa tidak memberikan jawaban."
            answer.save()
            
            # Recalculate session score
            session = answer.session
            session.score = calculate_score(session)
            session.save()
            
            return Response({
                'success': True,
                'score': 0,
                'feedback': answer.ai_feedback,
                'new_session_score': session.score
            })

        base_url = settings_obj.base_url.rstrip('/')
        url = f"{base_url}/chat/completions"
        headers = {
            "Authorization": f"Bearer {settings_obj.api_key}",
            "Content-Type": "application/json"
        }

        prompt = (
            f"Pertanyaan: {q.question_text}\n"
            f"Kunci Jawaban / Kriteria Penilaian: {q.explanation}\n"
            f"Jawaban Siswa: {student_text}\n"
            f"Poin Maksimal: {max_points}\n\n"
            f"Evaluasi jawaban siswa secara adil.\n"
            f"Kembalikan response dalam JSON dengan key 'score' (angka antara 0-{max_points}) dan 'feedback' (string masukan singkat dalam Bahasa Indonesia).\n"
            f"Contoh:\n"
            f"{{\n"
            f"  \"score\": 7.5,\n"
            f"  \"feedback\": \"Penjelasan sudah baik, tetapi kurang menyebutkan aspek Y.\"\n"
            f"}}"
        )

        payload = {
            "model": settings_obj.model_name,
            "messages": [
                {"role": "system", "content": "Anda adalah asisten koreksi ujian otomatis yang objektif dan mengembalikan format JSON murni."},
                {"role": "user", "content": prompt}
            ],
            "temperature": 0.2,
            "max_tokens": 1000
        }

        try:
            response = requests.post(url, json=payload, headers=headers, timeout=settings_obj.timeout)
            if response.status_code == 200:
                try:
                    resp_json = response.json()
                except Exception as json_err:
                    return Response({
                        'error': f"Gagal membaca JSON dari provider AI (Status 200 OK): {str(json_err)}. Response mentah: {response.text[:200]}"
                    }, status=status.HTTP_502_BAD_GATEWAY)
                
                try:
                    content = resp_json['choices'][0]['message']['content']
                except (KeyError, IndexError) as key_err:
                    return Response({
                        'error': f"Format respon AI tidak sesuai standard OpenAI: {str(key_err)}. Response: {str(resp_json)[:200]}"
                    }, status=status.HTTP_502_BAD_GATEWAY)
                
                try:
                    ai_result = clean_and_parse_json(content)
                except json.JSONDecodeError as parse_err:
                    return Response({
                        'error': f"Gagal mengurai hasil AI sebagai JSON: {str(parse_err)}. Teks dari AI: {content[:200]}"
                    }, status=status.HTTP_502_BAD_GATEWAY)
                
                score_val = ai_result.get('score', 0)
                score = Decimal(str(score_val))
                if score < 0:
                    score = Decimal('0')
                elif score > max_points:
                    score = Decimal(str(max_points))
                    
                feedback = ai_result.get('feedback', '')
                
                answer.points_earned = score
                answer.is_correct = score >= (Decimal(str(max_points)) / 2)
                answer.ai_feedback = feedback
                answer.save()
                
                # Recalculate session score
                session = answer.session
                session.score = calculate_score(session)
                session.save()
                
                return Response({
                    'success': True,
                    'score': float(score),
                    'feedback': feedback,
                    'new_session_score': session.score
                })
            else:
                return Response({'error': f"AI Provider returned HTTP {response.status_code}: {response.text[:200]}"}, status=status.HTTP_502_BAD_GATEWAY)
        except Exception as e:
            return Response({'error': f"Gagal menghubungi AI: {str(e)}"}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)



class TeacherExamSessionDetailView(generics.RetrieveDestroyAPIView):
    """Retrieve or delete detailed session answers for teachers/admins."""
    serializer_class = ExamSessionSerializer
    permission_classes = [IsAdminOrTeacher]
    queryset = ExamSession.objects.all()
    lookup_field = 'id'
    lookup_url_kwarg = 'session_id'

    def destroy(self, request, *args, **kwargs):
        session = self.get_object()
        # Also clean up orphan GuestParticipant if exists
        guest = session.guest_participant
        session.delete()
        if guest and not guest.exam_sessions.exists():
            guest.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class ManualGradeAnswerView(generics.GenericAPIView):
    """Manually grade a single student answer (essay)."""
    permission_classes = [IsAdminOrTeacher]

    def patch(self, request, answer_id):
        answer = generics.get_object_or_404(StudentAnswer, pk=answer_id)

        points = request.data.get('points_earned')
        feedback = request.data.get('feedback', '')

        if points is None:
            return Response({'error': 'points_earned wajib diisi.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            score = Decimal(str(points))
        except Exception:
            return Response({'error': 'Nilai poin tidak valid.'}, status=status.HTTP_400_BAD_REQUEST)

        max_points = answer.question.points
        if score < 0:
            score = Decimal('0')
        elif score > max_points:
            score = Decimal(str(max_points))

        answer.points_earned = score
        answer.is_correct = score >= (Decimal(str(max_points)) / 2)
        if feedback:
            answer.ai_feedback = f"[Manual] {feedback}"
        answer.save()

        # Recalculate session score
        session = answer.session
        session.score = calculate_score(session)
        session.save(update_fields=['score'])

        return Response({
            'id': answer.id,
            'points_earned': str(answer.points_earned),
            'is_correct': answer.is_correct,
            'ai_feedback': answer.ai_feedback,
            'new_session_score': str(session.score),
        })


class ExportSessionPDFView(generics.GenericAPIView):
    """Export an individual student's exam sheet/result to PDF."""
    permission_classes = [permissions.AllowAny]

    def get(self, request, session_id):
        from xhtml2pdf import pisa

        session = generics.get_object_or_404(ExamSession, pk=session_id)

        # Access check
        is_guest_allowed = False
        auth_header = request.META.get('HTTP_AUTHORIZATION', '')
        if auth_header.startswith('Guest '):
            token = auth_header[6:]
            try:
                guest = GuestParticipant.objects.get(access_token=token)
                if guest.is_token_valid and session.guest_participant == guest:
                    is_guest_allowed = True
            except GuestParticipant.DoesNotExist:
                pass

        is_user_allowed = False
        if request.user and request.user.is_authenticated:
            if request.user.role in ('admin', 'teacher') or session.student == request.user:
                is_user_allowed = True

        if not is_guest_allowed and not is_user_allowed:
            return Response({'error': 'Anda tidak memiliki akses ke lembar ujian ini.'}, status=status.HTTP_403_FORBIDDEN)

        exam = session.exam
        pg_questions = exam.questions.filter(question_type='multiple_choice').order_by('order')
        essay_questions = exam.questions.filter(question_type='essay').order_by('order')
        questions = list(pg_questions) + list(essay_questions)
        answers = {ans.question_id: ans for ans in session.answers.all().select_related('selected_choice')}

        # Formatting date and time
        hari_tanggal = '-'
        waktu = '-'
        if session.started_at:
            local_started_at = timezone.localtime(session.started_at)
            hari_tanggal = local_started_at.strftime('%A, %d %B %Y')
            waktu = local_started_at.strftime('%H:%M')
            
            day_translations = {
                'Monday': 'Senin', 'Tuesday': 'Selasa', 'Wednesday': 'Rabu',
                'Thursday': 'Kamis', 'Friday': 'Jumat', 'Saturday': 'Sabtu', 'Sunday': 'Minggu'
            }
            month_translations = {
                'January': 'Januari', 'February': 'Februari', 'March': 'Maret', 'April': 'April',
                'May': 'Mei', 'June': 'Juni', 'July': 'Juli', 'August': 'Agustus',
                'September': 'September', 'October': 'Oktober', 'November': 'November', 'December': 'Desember'
            }
            for eng, ind in day_translations.items():
                hari_tanggal = hari_tanggal.replace(eng, ind)
            for eng, ind in month_translations.items():
                hari_tanggal = hari_tanggal.replace(eng, ind)

            if session.submitted_at:
                local_submitted_at = timezone.localtime(session.submitted_at)
                waktu += f' - {local_submitted_at.strftime("%H:%M")}'
            else:
                waktu += ' - Selesai'

        nis = '-'
        kelas = '-'
        if session.student:
            profile = getattr(session.student, 'student_profile', None)
            if profile:
                nis = profile.nis
                kelas = profile.class_name
        elif session.guest_participant:
            nis = session.guest_participant.nis
            kelas = session.guest_participant.class_name

        total_earned = sum(float(ans.points_earned) for ans in answers.values())
        total_max = sum(q.points for q in questions)

        q_html = ''
        for idx, q in enumerate(questions, 1):
            ans = answers.get(q.id)
            points_earned = float(ans.points_earned) if ans else 0.0
            max_points = q.points
            
            if q.question_type == 'multiple_choice':
                choices = q.choices.all().order_by('order')
                choices_html = ''
                for c_idx, choice in enumerate(choices):
                    char = chr(65 + c_idx)
                    is_selected = ans and ans.selected_choice_id == choice.id
                    is_correct = choice.is_correct
                    
                    choice_style = ""
                    choice_label = ""
                    if is_selected:
                        choice_style = "font-weight: bold; color: #1e3a8a;"
                        if is_correct:
                            choice_label = " <span style='color: #16a34a; font-weight: bold;'>[Jawaban Siswa - Benar]</span>"
                        else:
                            choice_label = " <span style='color: #dc2626; font-weight: bold;'>[Jawaban Siswa - Salah]</span>"
                    elif is_correct:
                        choice_style = "color: #16a34a; font-weight: bold;"
                        choice_label = " <span style='color: #16a34a; font-weight: bold;'>[Kunci Jawaban]</span>"
                        
                    choices_html += f'''
                    <div class="choice-item" style="{choice_style}">
                        {char}. {choice.choice_text} {choice_label}
                    </div>
                    '''
                
                result_class = "correct" if ans and ans.is_correct else "incorrect"
                result_text = f"Benar (Skor: {points_earned:.1f} / {max_points:.1f})" if ans and ans.is_correct else f"Salah (Skor: {points_earned:.1f} / {max_points:.1f})"
                if not ans:
                    result_text = f"Tidak Dijawab (Skor: 0.0 / {max_points:.1f})"
                    
                q_html += f'''
                <div class="question-item">
                    <div class="question-text">{idx}. {q.question_text}</div>
                    <div class="choices-grid">
                        {choices_html}
                    </div>
                    <div class="answer-info {result_class}">
                        <strong>Status:</strong> {result_text}
                        {f'<div class="feedback-text">Feedback: {ans.ai_feedback}</div>' if ans and ans.ai_feedback else ''}
                    </div>
                </div>
                '''
            else:
                student_ans = ans.essay_answer if ans else ''
                student_ans_display = student_ans.replace('\n', '<br>') if student_ans else '<i>(Tidak Menjawab)</i>'
                
                result_class = "correct" if points_earned > (max_points / 2) else "incorrect"
                result_text = f"Skor: {points_earned:.1f} / {max_points:.1f}"
                
                feedback_str = ''
                if ans and ans.ai_feedback:
                    feedback_str = ans.ai_feedback.replace("[Manual] ", "")
                
                q_html += f'''
                <div class="question-item">
                    <div class="question-text">{idx}. {q.question_text}</div>
                    <div class="answer-info" style="background-color: #f8fafc; border-left: 3px solid #6366f1; color: #1e293b;">
                        <strong>Jawaban Siswa:</strong>
                        <p style="margin: 5px 0 0 0; font-family: monospace; font-size: 9.5pt;">{student_ans_display}</p>
                    </div>
                    <div class="answer-info {result_class}" style="margin-top: 5px;">
                        <strong>Hasil Penilaian:</strong> {result_text}
                        {f'<div class="feedback-text"><strong>Feedback/Analisis:</strong> {feedback_str}</div>' if feedback_str else ''}
                    </div>
                </div>
                '''

        tahun_pelajaran = "-"
        if session.started_at:
            year = int(session.started_at.strftime('%Y'))
            month = int(session.started_at.strftime('%m'))
            if month >= 7:
                tahun_pelajaran = f"{year}/{year + 1}"
            else:
                tahun_pelajaran = f"{year - 1}/{year}"

        html_content = f'''
        <html>
        <head>
            <style>
                @page {{
                    size: a4;
                    margin: 1.5cm;
                    margin-bottom: 2cm;
                    @frame footer {{
                        -pdf-frame-content: footerContent;
                        bottom: 1.2cm;
                        margin-left: 1.5cm;
                        margin-right: 1.5cm;
                        height: 1cm;
                    }}
                }}
                body {{
                    font-family: Arial, sans-serif;
                    font-size: 10pt;
                    line-height: 1.4;
                    color: #000000;
                }}
                .header {{
                    text-align: center;
                    margin-bottom: 5px;
                }}
                .header h1 {{
                    font-size: 13pt;
                    margin: 2px 0;
                    font-weight: bold;
                    text-transform: uppercase;
                }}
                .header h2 {{
                    font-size: 11pt;
                    margin: 2px 0;
                    font-weight: bold;
                    text-transform: uppercase;
                }}
                .header h3 {{
                    font-size: 10pt;
                    margin: 2px 0;
                    font-weight: normal;
                }}
                .double-line {{
                    border-top: 3px solid #000;
                    border-bottom: 1px solid #000;
                    height: 2px;
                    margin: 10px 0 15px 0;
                }}
                .doc-title {{
                    text-align: center;
                    font-size: 12pt;
                    font-weight: bold;
                    text-decoration: underline;
                    margin-bottom: 20px;
                    text-transform: uppercase;
                    letter-spacing: 1px;
                }}
                .meta-table {{
                    width: 100%;
                }}
                .meta-table td {{
                    padding: 3px 0;
                    vertical-align: top;
                    font-size: 9.5pt;
                }}
                .score-box {{
                    border: 2px solid #000000;
                    padding: 10px;
                    text-align: center;
                    width: 110px;
                    background-color: #ffffff;
                }}
                .score-box .score-title {{
                    font-size: 8pt;
                    font-weight: bold;
                    text-transform: uppercase;
                    margin-bottom: 3px;
                    letter-spacing: 0.5px;
                }}
                .score-box .score-value {{
                    font-size: 26pt;
                    font-weight: bold;
                    color: #000000;
                }}
                .question-section {{
                    margin-top: 10px;
                }}
                .question-item {{
                    margin-bottom: 18px;
                    page-break-inside: avoid;
                    border-bottom: 1px dashed #e2e8f0;
                    padding-bottom: 10px;
                }}
                .question-text {{
                    font-weight: bold;
                    margin-bottom: 6px;
                    font-size: 10pt;
                }}
                .choices-grid {{
                    margin-left: 15px;
                    margin-bottom: 8px;
                }}
                .choice-item {{
                    margin-bottom: 4px;
                    font-size: 9.5pt;
                }}
                .answer-info {{
                    padding: 6px 10px;
                    margin-left: 15px;
                    margin-top: 5px;
                    font-size: 9pt;
                    border-radius: 4px;
                }}
                .answer-info.correct {{
                    background-color: #e2f0d9;
                    border-left: 3px solid #385723;
                    color: #385723;
                }}
                .answer-info.incorrect {{
                    background-color: #fce4d6;
                    border-left: 3px solid #c65911;
                    color: #c65911;
                }}
                .feedback-text {{
                    font-style: italic;
                    margin-top: 4px;
                    color: #404040;
                }}
            </style>
        </head>
        <body>
            <div id="footerContent" style="text-align: right; font-size: 8pt; color: #555555; font-family: Arial, sans-serif;">
                Halaman <pdf:pagenumber>
            </div>
            
            <div class="header">
                <h1>PEMERINTAH KABUPATEN NEVUIAN</h1>
                <h2>DINAS PENDIDIKAN</h2>
                <h2>UPT DINAS PENDIDIKAN KECAMATAN NEVUIAN</h2>
                <h1 style="font-size: 15pt; margin-top: 6px;">{exam.title.upper()}</h1>
                <h3>Mata Pelajaran: {exam.subject} | Tahun Pelajaran {tahun_pelajaran}</h3>
            </div>
            
            <div class="double-line"></div>
            
            <div class="doc-title">LEMBAR HASIL UJIAN</div>
            
            <table style="width: 100%; border: none; margin-bottom: 15px;">
                <tr>
                    <td style="width: 70%; border: none; padding: 0; vertical-align: top;">
                        <table class="meta-table" style="border: none;">
                            <tr>
                                <td style="width: 120px; font-weight: bold;">MATA PELAJARAN</td>
                                <td style="width: 10px;">:</td>
                                <td style="font-weight: bold;">{exam.subject.upper()}</td>
                            </tr>
                            <tr>
                                <td style="font-weight: bold;">NAMA PESERTA</td>
                                <td>:</td>
                                <td>{session.participant_name.upper()}</td>
                            </tr>
                            <tr>
                                <td style="font-weight: bold;">NIS / USERNAME</td>
                                <td>:</td>
                                <td>{nis.upper()}</td>
                            </tr>
                            <tr>
                                <td style="font-weight: bold;">KELAS TARGET</td>
                                <td>:</td>
                                <td>{kelas.upper()}</td>
                            </tr>
                            <tr>
                                <td style="font-weight: bold;">HARI / TANGGAL</td>
                                <td>:</td>
                                <td>{hari_tanggal}</td>
                            </tr>
                            <tr>
                                <td style="font-weight: bold;">LAMA PENGERJAAN</td>
                                <td>:</td>
                                <td>{waktu}</td>
                            </tr>
                        </table>
                    </td>
                    <td style="width: 30%; border: none; padding: 0; text-align: right; vertical-align: top;">
                        <div class="score-box" style="display: inline-block;">
                            <div class="score-title">NILAI AKHIR</div>
                            <div class="score-value">{f"{float(session.score):.1f}" if session.score is not None else '-'}</div>
                            <div style="font-size: 7.5pt; color: #595959; margin-top: 3px;">Poin: {total_earned:.1f} / {total_max:.1f}</div>
                        </div>
                    </td>
                </tr>
            </table>
            
            <div style="border-top: 1px solid #000; margin-bottom: 15px;"></div>
            
            <div class="question-section">
                {q_html}
            </div>
        </body>
        </html>
        '''

        response = HttpResponse(content_type='application/pdf')
        filename = f'lembar_jawaban_{session.participant_name.replace(" ", "_")}_{exam.exam_code}.pdf'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        pisa.CreatePDF(html_content, dest=response)
        return response


class ExportSessionDocxView(generics.GenericAPIView):
    """Export an individual student's exam sheet/result to Word (.docx)."""
    permission_classes = [permissions.AllowAny]

    def get(self, request, session_id):
        import io
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import parse_xml

        session = generics.get_object_or_404(ExamSession, pk=session_id)

        # Access check
        is_guest_allowed = False
        auth_header = request.META.get('HTTP_AUTHORIZATION', '')
        if auth_header.startswith('Guest '):
            token = auth_header[6:]
            try:
                guest = GuestParticipant.objects.get(access_token=token)
                if guest.is_token_valid and session.guest_participant == guest:
                    is_guest_allowed = True
            except GuestParticipant.DoesNotExist:
                pass

        is_user_allowed = False
        if request.user and request.user.is_authenticated:
            if request.user.role in ('admin', 'teacher') or session.student == request.user:
                is_user_allowed = True

        if not is_guest_allowed and not is_user_allowed:
            return Response({'error': 'Anda tidak memiliki akses ke lembar ujian ini.'}, status=status.HTTP_403_FORBIDDEN)

        exam = session.exam
        pg_questions = exam.questions.filter(question_type='multiple_choice').order_by('order')
        essay_questions = exam.questions.filter(question_type='essay').order_by('order')
        questions = list(pg_questions) + list(essay_questions)
        answers = {ans.question_id: ans for ans in session.answers.all().select_related('selected_choice')}

        # Formatting date and time
        hari_tanggal = '-'
        waktu = '-'
        if session.started_at:
            local_started_at = timezone.localtime(session.started_at)
            hari_tanggal = local_started_at.strftime('%A, %d %B %Y')
            waktu = local_started_at.strftime('%H:%M')
            
            day_translations = {
                'Monday': 'Senin', 'Tuesday': 'Selasa', 'Wednesday': 'Rabu',
                'Thursday': 'Kamis', 'Friday': 'Jumat', 'Saturday': 'Sabtu', 'Sunday': 'Minggu'
            }
            month_translations = {
                'January': 'Januari', 'February': 'Februari', 'March': 'Maret', 'April': 'April',
                'May': 'Mei', 'June': 'Juni', 'July': 'Juli', 'August': 'Agustus',
                'September': 'September', 'October': 'Oktober', 'November': 'November', 'December': 'Desember'
            }
            for eng, ind in day_translations.items():
                hari_tanggal = hari_tanggal.replace(eng, ind)
            for eng, ind in month_translations.items():
                hari_tanggal = hari_tanggal.replace(eng, ind)

            if session.submitted_at:
                local_submitted_at = timezone.localtime(session.submitted_at)
                waktu += f' - {local_submitted_at.strftime("%H:%M")}'
            else:
                waktu += ' - Selesai'

        nis = '-'
        kelas = '-'
        if session.student:
            profile = getattr(session.student, 'student_profile', None)
            if profile:
                nis = profile.nis
                kelas = profile.class_name
        elif session.guest_participant:
            nis = session.guest_participant.nis
            kelas = session.guest_participant.class_name

        total_earned = sum(float(ans.points_earned) for ans in answers.values())
        total_max = sum(q.points for q in questions)

        tahun_pelajaran = "-"
        if session.started_at:
            year = int(session.started_at.strftime('%Y'))
            month = int(session.started_at.strftime('%m'))
            if month >= 7:
                tahun_pelajaran = f"{year}/{year + 1}"
            else:
                tahun_pelajaran = f"{year - 1}/{year}"

        doc = Document()
        
        # Set margins to 0.8 inches
        for s in doc.sections:
            s.top_margin = Inches(0.8)
            s.bottom_margin = Inches(0.8)
            s.left_margin = Inches(0.8)
            s.right_margin = Inches(0.8)

        # Center align text for Kop Surat
        p_kop = doc.add_paragraph()
        p_kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_kop.paragraph_format.space_after = Pt(2)
        
        r1 = p_kop.add_run("PEMERINTAH KABUPATEN NEVUIAN\n")
        r1.bold = True
        r1.font.size = Pt(13)
        r1.font.name = 'Arial'

        r2 = p_kop.add_run("DINAS PENDIDIKAN\n")
        r2.bold = True
        r2.font.size = Pt(11)
        r2.font.name = 'Arial'

        r3 = p_kop.add_run("UPT DINAS PENDIDIKAN KECAMATAN NEVUIAN\n")
        r3.bold = True
        r3.font.size = Pt(11)
        r3.font.name = 'Arial'

        r4 = p_kop.add_run(f"{exam.title.upper()}\n")
        r4.bold = True
        r4.font.size = Pt(14)
        r4.font.name = 'Arial'

        r5 = p_kop.add_run(f"Mata Pelajaran: {exam.subject} | Tahun Pelajaran {tahun_pelajaran}")
        r5.font.size = Pt(10)
        r5.font.name = 'Arial'

        # Add double line border under Kop Surat paragraph
        pPr = p_kop._p.get_or_add_pPr()
        pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                         r'  <w:bottom w:val="double" w:sz="18" w:space="8" w:color="000000"/>'
                         r'</w:pBdr>')
        pPr.append(pBdr)

        # Document Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(12)
        p_title.paragraph_format.space_after = Pt(12)
        r_title = p_title.add_run("LEMBAR HASIL UJIAN")
        r_title.bold = True
        r_title.font.size = Pt(12)
        r_title.font.underline = True
        r_title.font.name = 'Arial'

        # Metadata Table & Score Box Layout Table
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(4.8)
        table.columns[1].width = Inches(1.7)

        cell_left = table.cell(0, 0)
        cell_right = table.cell(0, 1)

        meta_data = [
            ("MATA PELAJARAN", exam.subject.upper()),
            ("NAMA PESERTA", session.participant_name.upper()),
            ("NIS / USERNAME", nis.upper()),
            ("KELAS TARGET", kelas.upper()),
            ("HARI / TANGGAL", hari_tanggal),
            ("LAMA PENGERJAAN", waktu),
        ]

        # Nested metadata table
        nested_table = cell_left.add_table(rows=len(meta_data), cols=3)
        nested_table.autofit = False
        nested_table.columns[0].width = Inches(1.6)
        nested_table.columns[1].width = Inches(0.2)
        nested_table.columns[2].width = Inches(3.0)

        for i, (label, val) in enumerate(meta_data):
            c0 = nested_table.cell(i, 0)
            c1 = nested_table.cell(i, 1)
            c2 = nested_table.cell(i, 2)
            
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(1)
            p0.paragraph_format.space_before = Pt(1)
            r0 = p0.add_run(label)
            r0.bold = True
            r0.font.name = 'Arial'
            r0.font.size = Pt(9)
            
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(1)
            p1.paragraph_format.space_before = Pt(1)
            r1 = p1.add_run(":")
            r1.font.name = 'Arial'
            r1.font.size = Pt(9)
            
            p2 = c2.paragraphs[0]
            p2.paragraph_format.space_after = Pt(1)
            p2.paragraph_format.space_before = Pt(1)
            r2 = p2.add_run(val)
            r2.font.name = 'Arial'
            r2.font.size = Pt(9)
            if label in ("MATA PELAJARAN", "NAMA PESERTA"):
                r2.bold = True

        # Score Box formatting
        tcPr = cell_right._tc.get_or_add_tcPr()
        tcBorders = parse_xml(r'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                              r'  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
                              r'  <w:left w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
                              r'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
                              r'  <w:right w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
                              r'</w:tcBorders>')
        tcPr.append(tcBorders)

        tcMar = parse_xml(r'<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                          r'  <w:top w:w="120" w:type="dxa"/>'
                          r'  <w:left w:w="120" w:type="dxa"/>'
                          r'  <w:bottom w:w="120" w:type="dxa"/>'
                          r'  <w:right w:w="120" w:type="dxa"/>'
                          r'</w:tcMar>')
        tcPr.append(tcMar)

        p_score_lbl = cell_right.paragraphs[0]
        p_score_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_score_lbl.paragraph_format.space_after = Pt(2)
        r_lbl = p_score_lbl.add_run("NILAI AKHIR")
        r_lbl.bold = True
        r_lbl.font.size = Pt(8.5)
        r_lbl.font.name = 'Arial'

        p_score_val = cell_right.add_paragraph()
        p_score_val.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_score_val.paragraph_format.space_before = Pt(2)
        p_score_val.paragraph_format.space_after = Pt(2)
        score_str = f"{float(session.score):.1f}" if session.score is not None else '-'
        r_val = p_score_val.add_run(score_str)
        r_val.bold = True
        r_val.font.size = Pt(26)
        r_val.font.name = 'Arial'

        p_score_sub = cell_right.add_paragraph()
        p_score_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_score_sub.paragraph_format.space_before = Pt(2)
        p_score_sub.paragraph_format.space_after = Pt(0)
        r_sub = p_score_sub.add_run(f"Poin: {total_earned:.1f} / {total_max:.1f}")
        r_sub.font.size = Pt(7.5)
        r_sub.font.name = 'Arial'
        r_sub.font.color.rgb = RGBColor(89, 89, 89)

        # Separator line
        p_sep = doc.add_paragraph()
        p_sep.paragraph_format.space_before = Pt(12)
        p_sep.paragraph_format.space_after = Pt(12)
        pPr_sep = p_sep._p.get_or_add_pPr()
        pBdr_sep = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                             r'  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/>'
                             r'</w:pBdr>')
        pPr_sep.append(pBdr_sep)

        # Questions list
        for idx, q in enumerate(questions, 1):
            ans = answers.get(q.id)
            points_earned = float(ans.points_earned) if ans else 0.0
            max_points = q.points
            
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(8)
            p_q.paragraph_format.space_after = Pt(4)
            p_q.paragraph_format.keep_with_next = True
            
            r_q = p_q.add_run(f"{idx}. {q.question_text}")
            r_q.bold = True
            r_q.font.name = 'Arial'
            r_q.font.size = Pt(10.5)

            if q.question_type == 'multiple_choice':
                choices = q.choices.all().order_by('order')
                for c_idx, choice in enumerate(choices):
                    char = chr(65 + c_idx)
                    is_selected = ans and ans.selected_choice_id == choice.id
                    is_correct = choice.is_correct
                    
                    p_c = doc.add_paragraph()
                    p_c.paragraph_format.left_indent = Inches(0.25)
                    p_c.paragraph_format.space_before = Pt(1)
                    p_c.paragraph_format.space_after = Pt(1)
                    
                    r_char = p_c.add_run(f"{char}. {choice.choice_text}")
                    r_char.font.name = 'Arial'
                    r_char.font.size = Pt(9.5)
                    
                    if is_selected:
                        r_char.bold = True
                        if is_correct:
                            r_lbl = p_c.add_run(" [Jawaban Siswa - Benar]")
                            r_lbl.bold = True
                            r_lbl.font.color.rgb = RGBColor(22, 163, 74)
                        else:
                            r_lbl = p_c.add_run(" [Jawaban Siswa - Salah]")
                            r_lbl.bold = True
                            r_lbl.font.color.rgb = RGBColor(220, 38, 38)
                    elif is_correct:
                        r_char.font.color.rgb = RGBColor(22, 163, 74)
                        r_lbl = p_c.add_run(" [Kunci Jawaban]")
                        r_lbl.bold = True
                        r_lbl.font.color.rgb = RGBColor(22, 163, 74)

                # Shaded alert box for status
                p_stat = doc.add_paragraph()
                p_stat.paragraph_format.left_indent = Inches(0.25)
                p_stat.paragraph_format.space_before = Pt(4)
                p_stat.paragraph_format.space_after = Pt(8)
                
                result_text = f"Benar (Skor: {points_earned:.1f} / {max_points:.1f})" if ans and ans.is_correct else f"Salah (Skor: {points_earned:.1f} / {max_points:.1f})"
                if not ans:
                    result_text = f"Tidak Dijawab (Skor: 0.0 / {max_points:.1f})"
                    
                r_stat_lbl = p_stat.add_run("Status: ")
                r_stat_lbl.bold = True
                r_stat_lbl.font.name = 'Arial'
                r_stat_lbl.font.size = Pt(9)
                
                r_stat_val = p_stat.add_run(result_text)
                r_stat_val.font.name = 'Arial'
                r_stat_val.font.size = Pt(9)
                if ans and ans.is_correct:
                    r_stat_val.font.color.rgb = RGBColor(56, 87, 35)
                else:
                    r_stat_val.font.color.rgb = RGBColor(198, 89, 17)
                    
                if ans and ans.ai_feedback:
                    p_stat.add_run("\nFeedback: ")
                    r_fb = p_stat.add_run(ans.ai_feedback)
                    r_fb.italic = True
                    r_fb.font.name = 'Arial'
                    r_fb.font.size = Pt(8.5)
                    r_fb.font.color.rgb = RGBColor(64, 64, 64)
                    
                pPr = p_stat._p.get_or_add_pPr()
                pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                                 r'  <w:left w:val="single" w:sz="24" w:space="4" w:color="' + ('385723' if ans and ans.is_correct else 'C65911') + r'"/>'
                                 r'</w:pBdr>')
                pPr.append(pBdr)
                shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="' + ('E2F0D9' if ans and ans.is_correct else 'FCE4D6') + r'"/>')
                pPr.append(shd)

            else:
                student_ans = ans.essay_answer if ans else ''
                student_ans_display = student_ans if student_ans else '(Tidak Menjawab)'
                
                p_ans_title = doc.add_paragraph()
                p_ans_title.paragraph_format.left_indent = Inches(0.25)
                p_ans_title.paragraph_format.space_before = Pt(4)
                p_ans_title.paragraph_format.space_after = Pt(2)
                r_ans_title = p_ans_title.add_run("Jawaban Siswa:")
                r_ans_title.bold = True
                r_ans_title.font.name = 'Arial'
                r_ans_title.font.size = Pt(9)
                
                p_ans = doc.add_paragraph()
                p_ans.paragraph_format.left_indent = Inches(0.25)
                p_ans.paragraph_format.space_before = Pt(0)
                p_ans.paragraph_format.space_after = Pt(6)
                r_ans_val = p_ans.add_run(student_ans_display)
                r_ans_val.font.name = 'Courier New'
                r_ans_val.font.size = Pt(9.5)
                
                pPr_ans = p_ans._p.get_or_add_pPr()
                pBdr_ans = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                                     r'  <w:left w:val="single" w:sz="24" w:space="4" w:color="6366F1"/>'
                                     r'</w:pBdr>')
                pPr_ans.append(pBdr_ans)
                shd_ans = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F8FAFC"/>')
                pPr_ans.append(shd_ans)
                
                p_grad = doc.add_paragraph()
                p_grad.paragraph_format.left_indent = Inches(0.25)
                p_grad.paragraph_format.space_before = Pt(4)
                p_grad.paragraph_format.space_after = Pt(8)
                
                result_text = f"Skor: {points_earned:.1f} / {max_points:.1f}"
                is_pass = points_earned >= (max_points / 2)
                
                r_grad_lbl = p_grad.add_run("Hasil Penilaian: ")
                r_grad_lbl.bold = True
                r_grad_lbl.font.name = 'Arial'
                r_grad_lbl.font.size = Pt(9)
                
                r_grad_val = p_grad.add_run(result_text)
                r_grad_val.font.name = 'Arial'
                r_grad_val.font.size = Pt(9)
                r_grad_val.font.color.rgb = RGBColor(56, 87, 35) if is_pass else RGBColor(198, 89, 17)
                
                feedback_str = ''
                if ans and ans.ai_feedback:
                    feedback_str = ans.ai_feedback.replace("[Manual] ", "")
                    
                if feedback_str:
                    p_grad.add_run("\nFeedback/Analisis: ")
                    r_fb = p_grad.add_run(feedback_str)
                    r_fb.italic = True
                    r_fb.font.name = 'Arial'
                    r_fb.font.size = Pt(8.5)
                    r_fb.font.color.rgb = RGBColor(64, 64, 64)
                    
                pPr_grad = p_grad._p.get_or_add_pPr()
                pBdr_grad = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                                      r'  <w:left w:val="single" w:sz="24" w:space="4" w:color="' + ('385723' if is_pass else 'C65911') + r'"/>'
                                      r'</w:pBdr>')
                pPr_grad.append(pBdr_grad)
                shd_grad = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="' + ('E2F0D9' if is_pass else 'FCE4D6') + r'"/>')
                pPr_grad.append(shd_grad)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        filename = f'lembar_jawaban_{session.participant_name.replace(" ", "_")}_{exam.exam_code}.docx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class ExportQuestionsDocxView(generics.GenericAPIView):
    """Export exam questions to Word (.docx) for print/distribution."""
    permission_classes = [IsAdminOrTeacher]

    def get(self, request, exam_id):
        import io
        from django.utils import timezone
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import parse_xml

        exam = generics.get_object_or_404(Exam, pk=exam_id)
        pg_questions = exam.questions.filter(question_type='multiple_choice').order_by('order')
        essay_questions = exam.questions.filter(question_type='essay').order_by('order')
        questions = list(pg_questions) + list(essay_questions)

        tahun_pelajaran = "-"
        now = timezone.now()
        year = int(now.strftime('%Y'))
        month = int(now.strftime('%m'))
        if month >= 7:
            tahun_pelajaran = f"{year}/{year + 1}"
        else:
            tahun_pelajaran = f"{year - 1}/{year}"

        doc = Document()
        
        # Set margins to 0.8 inches
        for s in doc.sections:
            s.top_margin = Inches(0.8)
            s.bottom_margin = Inches(0.8)
            s.left_margin = Inches(0.8)
            s.right_margin = Inches(0.8)

        # Center align text for Kop Surat
        p_kop = doc.add_paragraph()
        p_kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_kop.paragraph_format.space_after = Pt(2)
        
        r1 = p_kop.add_run("PEMERINTAH KABUPATEN NEVUIAN\n")
        r1.bold = True
        r1.font.size = Pt(13)
        r1.font.name = 'Arial'

        r2 = p_kop.add_run("DINAS PENDIDIKAN\n")
        r2.bold = True
        r2.font.size = Pt(11)
        r2.font.name = 'Arial'

        r3 = p_kop.add_run("UPT DINAS PENDIDIKAN KECAMATAN NEVUIAN\n")
        r3.bold = True
        r3.font.size = Pt(11)
        r3.font.name = 'Arial'

        r4 = p_kop.add_run(f"{exam.title.upper()}\n")
        r4.bold = True
        r4.font.size = Pt(14)
        r4.font.name = 'Arial'

        r5 = p_kop.add_run(f"TAHUN PELAJARAN {tahun_pelajaran}")
        r5.font.size = Pt(10)
        r5.font.name = 'Arial'

        # Add double line border under Kop Surat paragraph
        pPr = p_kop._p.get_or_add_pPr()
        pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                         r'  <w:bottom w:val="double" w:sz="18" w:space="8" w:color="000000"/>'
                         r'</w:pBdr>')
        pPr.append(pBdr)

        # Document Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(12)
        p_title.paragraph_format.space_after = Pt(12)
        r_title = p_title.add_run("LEMBAR SOAL")
        r_title.bold = True
        r_title.font.size = Pt(13)
        r_title.font.name = 'Arial'

        # Metadata Table
        meta_table = doc.add_table(rows=4, cols=3)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False
        meta_table.columns[0].width = Inches(2.2)
        meta_table.columns[1].width = Inches(0.2)
        meta_table.columns[2].width = Inches(4.1)

        meta_rows = [
            ("MATA PELAJARAN", exam.subject.upper()),
            ("KELAS / TARGET", exam.class_target.upper()),
            ("WAKTU PENGERJAAN", f"{exam.duration_minutes} MENIT"),
            ("HARI / TANGGAL", "............................................."),
        ]

        for i, (label, val) in enumerate(meta_rows):
            c0 = meta_table.cell(i, 0)
            c1 = meta_table.cell(i, 1)
            c2 = meta_table.cell(i, 2)
            
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(2)
            p0.paragraph_format.space_before = Pt(2)
            r0 = p0.add_run(label)
            r0.bold = True
            r0.font.name = 'Arial'
            r0.font.size = Pt(9.5)
            
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(2)
            p1.paragraph_format.space_before = Pt(2)
            r1 = p1.add_run(":")
            r1.font.name = 'Arial'
            r1.font.size = Pt(9.5)
            
            p2 = c2.paragraphs[0]
            p2.paragraph_format.space_after = Pt(2)
            p2.paragraph_format.space_before = Pt(2)
            r2 = p2.add_run(val)
            r2.font.name = 'Arial'
            r2.font.size = Pt(9.5)
            if label != "HARI / TANGGAL":
                r2.bold = True

        # Separator line
        p_sep = doc.add_paragraph()
        p_sep.paragraph_format.space_before = Pt(12)
        p_sep.paragraph_format.space_after = Pt(12)
        pPr_sep = p_sep._p.get_or_add_pPr()
        pBdr_sep = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                             r'  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/>'
                             r'</w:pBdr>')
        pPr_sep.append(pBdr_sep)

        # Petunjuk Umum Section
        p_petunjuk_hdr = doc.add_paragraph()
        p_petunjuk_hdr.paragraph_format.space_before = Pt(4)
        p_petunjuk_hdr.paragraph_format.space_after = Pt(2)
        r_ph = p_petunjuk_hdr.add_run("PETUNJUK UMUM:")
        r_ph.bold = True
        r_ph.font.name = 'Arial'
        r_ph.font.size = Pt(9.5)

        petunjuk_list = [
            "Berdoalah sebelum mulai mengerjakan soal.",
            "Tulislah identitas Anda pada lembar jawaban yang tersedia.",
            "Bacalah soal-soal dengan teliti sebelum Anda menjawabnya.",
            "Dahulukan menjawab soal-soal yang Anda anggap mudah.",
            "Periksalah pekerjaan Anda sebelum diserahkan kepada pengawas.",
        ]
        for idx_p, petunjuk in enumerate(petunjuk_list, 1):
            p_pet = doc.add_paragraph()
            p_pet.paragraph_format.left_indent = Inches(0.2)
            p_pet.paragraph_format.space_after = Pt(1)
            p_pet.paragraph_format.space_before = Pt(1)
            r_p = p_pet.add_run(f"{idx_p}. {petunjuk}")
            r_p.font.name = 'Arial'
            r_p.font.size = Pt(9)
            r_p.font.color.rgb = RGBColor(64, 64, 64)

        # Separator line after instructions
        p_sep2 = doc.add_paragraph()
        p_sep2.paragraph_format.space_before = Pt(8)
        p_sep2.paragraph_format.space_after = Pt(16)
        pPr_sep2 = p_sep2._p.get_or_add_pPr()
        pBdr_sep2 = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                              r'  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/>'
                              r'</w:pBdr>')
        pPr_sep2.append(pBdr_sep2)

        # Questions listing
        for idx, q in enumerate(questions, 1):
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(8)
            p_q.paragraph_format.space_after = Pt(4)
            p_q.paragraph_format.keep_with_next = True
            
            r_q = p_q.add_run(f"{idx}. {q.question_text}")
            r_q.bold = True
            r_q.font.name = 'Arial'
            r_q.font.size = Pt(10.5)

            if q.question_type == 'multiple_choice':
                choices = q.choices.all().order_by('order')
                for c_idx, choice in enumerate(choices):
                    char = chr(65 + c_idx)
                    p_c = doc.add_paragraph()
                    p_c.paragraph_format.left_indent = Inches(0.25)
                    p_c.paragraph_format.space_before = Pt(1)
                    p_c.paragraph_format.space_after = Pt(1)
                    
                    r_char = p_c.add_run(f"{char}. {choice.choice_text}")
                    r_char.font.name = 'Arial'
                    r_char.font.size = Pt(9.5)
            else:
                # Essay question, leave space for answers
                p_ans_space = doc.add_paragraph()
                p_ans_space.paragraph_format.left_indent = Inches(0.25)
                p_ans_space.paragraph_format.space_before = Pt(6)
                p_ans_space.paragraph_format.space_after = Pt(6)
                
                # Draw lines for student to write
                r_lines = p_ans_space.add_run(
                    "Jawaban:\n"
                    "_________________________________________________________________________________\n"
                    "_________________________________________________________________________________\n"
                    "_________________________________________________________________________________"
                )
                r_lines.font.name = 'Arial'
                r_lines.font.size = Pt(9.5)
                r_lines.font.color.rgb = RGBColor(128, 128, 128)

        # Save to memory stream
        output = io.BytesIO()
        doc.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        filename = f'lembar_soal_{exam.title.replace(" ", "_")}_{exam.exam_code}.docx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


